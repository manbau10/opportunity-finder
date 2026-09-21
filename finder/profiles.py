"""CV ingestion and structured profile extraction."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import re
import uuid
from collections import Counter

from . import store
from .auth import now
from .domains import DOMAIN_LABELS, TARGET_TITLES, primary_domain, ranked_domains

MAX_CV_BYTES = 6 * 1024 * 1024
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "md"}

STOPWORDS = {
    "about", "after", "again", "also", "among", "and", "are", "been", "being",
    "between", "both", "can", "from", "have", "into", "more", "most", "other",
    "our", "over", "such", "than", "that", "the", "their", "there", "these",
    "they", "this", "through", "using", "was", "were", "which", "while", "with",
    "work", "year", "years", "your", "university", "research", "experience",
}

PHRASES = [
    "construction management", "construction engineering", "civil engineering",
    "infrastructure engineering", "infrastructure management", "project management",
    "water infrastructure", "water resources", "asset management", "built environment",
    "building information modelling", "digital twin", "artificial intelligence",
    "machine learning", "data analytics", "sustainability", "risk management",
    "contract management", "quantity surveying", "structural engineering",
    "environmental engineering", "transport infrastructure", "stakeholder management",
    "programme management", "cost management", "health and safety", "procurement",
]


def _extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text(data: bytes, filename: str) -> tuple[str, str]:
    if not data or len(data) > MAX_CV_BYTES:
        raise ValueError("CV files must be between 1 byte and 6 MB.")
    ext = _extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Upload a PDF, DOCX, TXT or Markdown CV.")
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        mime = "application/pdf"
    elif ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(chunks)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        text = data.decode("utf-8", errors="replace")
        mime = "text/plain"
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 120:
        raise ValueError("Very little readable text was found in this CV.")
    return text[:250000], mime


def build_profile(text: str, kind: str) -> dict:
    lower = text.lower()
    phrase_hits = [p for p in PHRASES if p in lower]
    words = re.findall(r"[a-z][a-z0-9+#.-]{2,}", lower)
    counts = Counter(w for w in words if w not in STOPWORDS and not w.isdigit())
    keywords = phrase_hits + [w for w, _ in counts.most_common(45) if w not in phrase_hits]
    emails = re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    phones = re.findall(r"(?:\+?\d[\d ()-]{7,}\d)", text)
    degree = "PhD" if re.search(r"\b(ph\.?d|doctor of philosophy)\b", lower) else (
        "Master's" if re.search(r"\b(m\.?sc|m\.?eng|master)\b", lower) else "")
    domain = primary_domain(text)
    subject_titles = TARGET_TITLES.get(domain, [])
    target_titles = (
        ["assistant professor", "lecturer", "research fellow", "postdoctoral fellow"]
        if kind == "academic" else subject_titles
    )
    return {
        "kind": kind,
        "keywords": list(dict.fromkeys(keywords))[:55],
        "target_titles": target_titles,
        "primary_domain": domain,
        "domain_label": DOMAIN_LABELS.get(domain, "General / multidisciplinary"),
        "domains": ranked_domains(text)[:5],
        "degree": degree,
        "email": emails[0] if emails else "",
        "phone": phones[0].strip() if phones else "",
        "word_count": len(words),
    }


def save_profile(user_id: str, kind: str, filename: str, data: bytes) -> dict:
    if kind not in ("academic", "industry"):
        raise ValueError("Unknown CV type.")
    text, mime = extract_text(data, filename)
    profile = build_profile(text, kind)
    stamp = now()
    digest = hashlib.sha256(data).hexdigest()
    conn = store.connect()
    try:
        old = conn.execute(
            "SELECT id,created_at FROM user_profiles WHERE user_id=? AND kind=?",
            (user_id, kind),
        ).fetchone()
        if old:
            profile_id, created = old["id"], old["created_at"]
            conn.execute(
                """UPDATE user_profiles SET cv_filename=?,cv_mime=?,cv_sha256=?,
                   cv_blob=?,cv_text=?,profile_json=?,updated_at=? WHERE id=? AND user_id=?""",
                (filename[:240], mime, digest, base64.b64encode(data).decode("ascii"), text,
                 json.dumps(profile), stamp, profile_id, user_id),
            )
        else:
            profile_id, created = uuid.uuid4().hex, stamp
            conn.execute(
                """INSERT INTO user_profiles
                   (id,user_id,kind,cv_filename,cv_mime,cv_sha256,cv_blob,cv_text,
                    profile_json,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (profile_id, user_id, kind, filename[:240], mime, digest,
                 base64.b64encode(data).decode("ascii"), text, json.dumps(profile),
                 created, stamp),
            )
        conn.commit()
    finally:
        conn.close()
    return {"id": profile_id, "kind": kind, "filename": filename, "profile": profile,
            "created_at": created, "updated_at": stamp}


def get_profile(user_id: str, kind: str, include_text: bool = False) -> dict | None:
    # Always fetch the extracted text so profiles created before occupational
    # classification was introduced can be upgraded transparently.
    fields = "id,user_id,kind,cv_filename,cv_mime,cv_sha256,cv_text,profile_json,created_at,updated_at"
    conn = store.connect()
    try:
        row = conn.execute(
            f"SELECT {fields} FROM user_profiles WHERE user_id=? AND kind=?",
            (user_id, kind),
        ).fetchone()
    finally:
        conn.close()
    if not row:
        return None
    result = dict(row)
    result["profile"] = json.loads(result.pop("profile_json") or "{}")
    if not result["profile"].get("primary_domain"):
        result["profile"] = build_profile(result.get("cv_text") or "", kind)
        conn = store.connect()
        try:
            conn.execute("UPDATE user_profiles SET profile_json=?,updated_at=? WHERE id=? AND user_id=?",
                         (json.dumps(result["profile"]), now(), result["id"], user_id))
            conn.commit()
        finally:
            conn.close()
    if not include_text:
        result.pop("cv_text", None)
    return result


def profile_summaries(user_id: str) -> dict:
    return {kind: get_profile(user_id, kind) for kind in ("academic", "industry")}
