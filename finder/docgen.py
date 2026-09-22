"""Deterministic, professionally formatted application-document builders."""

from __future__ import annotations

import copy
import datetime as dt
import io
import pathlib
import re
import textwrap
from typing import Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate
import reportlab


BLACK = RGBColor(0, 0, 0)


def _pdf_fonts(sample: str = "") -> tuple[str, str]:
    cid_choices = (
        (r"[\u3040-\u30ff]", "HeiseiMin-W3"),
        (r"[\uac00-\ud7af]", "HYSMyeongJo-Medium"),
        (r"[\u3400-\u9fff]", "STSong-Light"),
    )
    for pattern, font_name in cid_choices:
        if re.search(pattern, sample):
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(UnicodeCIDFont(font_name))
            return font_name, font_name
    candidates = [
        pathlib.Path(reportlab.__file__).parent / "fonts" / "Vera.ttf",
        pathlib.Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        pathlib.Path("C:/Windows/Fonts/arial.ttf"),
    ]
    bold_candidates = [
        pathlib.Path(reportlab.__file__).parent / "fonts" / "VeraBd.ttf",
        pathlib.Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        pathlib.Path("C:/Windows/Fonts/arialbd.ttf"),
    ]
    regular = next((path for path in candidates if path.exists()), None)
    bold = next((path for path in bold_candidates if path.exists()), None)
    if regular and bold:
        if "ApplicationSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("ApplicationSans", str(regular)))
            pdfmetrics.registerFont(TTFont("ApplicationSansBold", str(bold)))
        return "ApplicationSans", "ApplicationSansBold"
    return "Helvetica", "Helvetica-Bold"


def _set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def base_doc(title: str, subtitle: str = "", *, page_system: str = "letter",
             show_title: bool = True) -> Document:
    doc = Document()
    section = doc.sections[0]
    if page_system == "a4":
        section.page_width, section.page_height = Mm(210), Mm(297)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(4)
    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(10.5)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)

    if "Application Metadata" not in styles:
        meta = styles.add_style("Application Metadata", WD_STYLE_TYPE.PARAGRAPH)
        meta.font.name = "Aptos"
        meta.font.size = Pt(9.5)
        meta.font.color.rgb = RGBColor(70, 70, 70)
        meta.paragraph_format.space_after = Pt(2)

    if show_title:
        title_p = doc.add_paragraph(style="Title")
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(title_p.add_run(title), "Aptos Display", 18, True)
        if subtitle:
            p = doc.add_paragraph(style="Subtitle")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(subtitle)
    footer = section.footer.paragraphs[0]
    footer.style = styles["Application Metadata"]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(title + "  |  ")
    _add_page_number(footer)
    return doc


def _add_rich_text(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _append_content(doc: Document, content: str) -> None:
    lines = (content or "").replace("\r\n", "\n").split("\n")
    pending: list[str] = []

    def flush() -> None:
        if not pending:
            return
        text = " ".join(x.strip() for x in pending if x.strip())
        pending.clear()
        if text:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_rich_text(p, text)

    for raw in lines:
        line = raw.strip()
        if not line:
            flush()
        elif line.startswith("### "):
            flush()
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            flush()
            doc.add_heading(line[3:].strip(), level=1)
        elif re.match(r"^[-*•]\s+", line):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            _add_rich_text(p, re.sub(r"^[-*•]\s+", "", line))
        elif re.match(r"^\d+[.)]\s+", line):
            flush()
            p = doc.add_paragraph(style="List Number")
            _add_rich_text(p, re.sub(r"^\d+[.)]\s+", "", line))
        else:
            pending.append(line)
    flush()


def docx_bytes(title: str, content: str, subtitle: str = "", *,
               page_system: str = "letter") -> bytes:
    doc = base_doc(title, subtitle, page_system=page_system)
    _append_content(doc, content)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def application_docx(title: str, content: str, opportunity: dict, applicant: dict,
                     document_type: str, *, page_system: str = "letter") -> bytes:
    role = opportunity.get("title") or "Application"
    org = opportunity.get("org") or "Employer"
    name = applicant.get("name") or "Applicant"
    is_letter = document_type == "cover_letter"
    doc = base_doc(title, "", page_system=page_system, show_title=not is_letter)
    if is_letter:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _set_font(p.add_run(name), "Aptos Display", 15, True)
        contact = "  |  ".join(x for x in (applicant.get("email"), applicant.get("phone")) if x)
        if contact:
            doc.add_paragraph(contact, style="Application Metadata")
        doc.add_paragraph(dt.date.today().strftime("%d %B %Y"), style="Application Metadata")
        doc.add_paragraph(org)
        if opportunity.get("department"):
            doc.add_paragraph(opportunity["department"])
        doc.add_paragraph()
        subject = doc.add_paragraph()
        subject.paragraph_format.space_after = Pt(10)
        _set_font(subject.add_run(f"Application for {role}"), "Aptos", 11, True)
    else:
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.style = doc.styles["Application Metadata"]
        meta.add_run(f"{name}  |  {role}  |  {org}")
    _append_content(doc, content)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _shading(fill: str):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    return node


def requirements_doc(requirements: dict, opportunity: dict) -> bytes:
    doc = base_doc(
        "Application Requirements and Evidence Plan",
        f"{opportunity.get('title', '')} at {opportunity.get('org', '')}",
        page_system=page_system_for(opportunity),
    )
    doc.add_heading("Submission checklist", level=1)
    for item in requirements.get("submission_items", []):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item.get("label", "Document")).bold = True
        detail = item.get("detail") or "Required"
        if item.get("limit"):
            detail += f"; limit: {item['limit']}"
        p.add_run(f" - {detail}")
    doc.add_heading("Selection criteria and evidence", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, label in enumerate(("Criterion", "Priority", "Evidence to use")):
        cell = table.rows[0].cells[idx]
        cell.text = label
        cell._tc.get_or_add_tcPr().append(_shading("1F4E79"))
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
        _set_cell_margins(cell)
    _set_repeat_table_header(table.rows[0])
    for row_no, criterion in enumerate(requirements.get("criteria", [])):
        cells = table.add_row().cells
        values = (
            criterion.get("criterion", ""), criterion.get("priority", ""),
            criterion.get("evidence", "Needs applicant verification"),
        )
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            _set_cell_margins(cells[idx])
            if row_no % 2:
                cells[idx]._tc.get_or_add_tcPr().append(_shading("F2F6FA"))
    doc.add_heading("Important gaps and checks", level=1)
    for gap in requirements.get("gaps", []) or ["Verify every claim before submission."]:
        doc.add_paragraph(gap, style="List Bullet")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def sources_doc(sources: list[dict], opportunity: dict) -> bytes:
    doc = base_doc(
        "Research Sources Used for the Application",
        f"{opportunity.get('title', 'Role')} at {opportunity.get('org', 'Employer')}",
        page_system=page_system_for(opportunity),
    )
    doc.add_paragraph(
        "These sources were collected for application research. Check the live pages before "
        "submitting because course, staff and organisational information can change."
    )
    for n, source in enumerate(sources, 1):
        doc.add_heading(f"{n}  {source.get('title') or 'Untitled source'}", level=1)
        meta = doc.add_paragraph(style="Application Metadata")
        meta.add_run(f"Type: {source.get('kind', 'web page')}  |  Retrieved: {source.get('retrieved_at', '')}")
        if source.get("url"):
            p = doc.add_paragraph()
            add_hyperlink(p, source["url"], source["url"])
        if source.get("query"):
            doc.add_paragraph(f"Search query: {source['query']}", style="Application Metadata")
        doc.add_paragraph(source.get("snippet") or "No extract was available.")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def job_doc(opportunity: dict, full_text: str = "") -> bytes:
    body = (
        f"## Role details\n\n**Title:** {opportunity.get('title','')}\n\n"
        f"**Employer:** {opportunity.get('org','')}\n\n"
        f"**Department:** {opportunity.get('department','')}\n\n"
        f"**Location:** {opportunity.get('location','')}\n\n"
        f"**Deadline:** {opportunity.get('deadline') or 'Not stated'}\n\n"
        f"**Original advert:** {opportunity.get('url','')}\n\n"
        f"## Advert text\n\n{full_text or opportunity.get('description','')}"
    )
    return docx_bytes("Job Advertisement Snapshot", body, page_system=page_system_for(opportunity))


def _iter_paragraphs(doc: Document) -> Iterable:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _replace_paragraph_text(paragraph, text: str) -> None:
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(text)
        return
    target = next((r for r in runs if r.text.strip()), runs[0])
    target.text = text
    for run in runs:
        if run is not target:
            run.text = ""


def _insert_after(paragraph, text: str, style: str = "List Bullet") -> None:
    new_p = copy.deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag != qn("w:pPr"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    inserted = Paragraph(new_p, paragraph._parent)
    try:
        inserted.style = style
    except KeyError:
        pass
    inserted.add_run(text)


def tailor_cv_docx(original: bytes, edits: dict) -> tuple[bytes, dict]:
    """Apply conservative edits to a copy while retaining the uploaded DOCX design."""
    doc = Document(io.BytesIO(original))
    paragraphs = list(_iter_paragraphs(doc))
    replacements_applied = additions_applied = 0
    for edit in edits.get("replacements", [])[:20]:
        find = " ".join(str(edit.get("find", "")).split())
        replacement = str(edit.get("replacement", "")).strip()
        if not find or not replacement:
            continue
        for paragraph in paragraphs:
            if " ".join(paragraph.text.split()) == find:
                _replace_paragraph_text(paragraph, replacement)
                replacements_applied += 1
                break
    for edit in edits.get("additions", [])[:16]:
        anchor = " ".join(str(edit.get("after", "")).split())
        text = str(edit.get("text", "")).strip()
        if not anchor or not text:
            continue
        for paragraph in paragraphs:
            if " ".join(paragraph.text.split()) == anchor:
                _insert_after(paragraph, text, edit.get("style") or "List Bullet")
                additions_applied += 1
                break
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue(), {
        "replacements_requested": len(edits.get("replacements", [])),
        "replacements_applied": replacements_applied,
        "additions_requested": len(edits.get("additions", [])),
        "additions_applied": additions_applied,
    }


def webpage_pdf(title: str, url: str, text: str, retrieved_at: str,
                *, page_system: str = "letter") -> bytes:
    output = io.BytesIO()
    pagesize = A4 if page_system == "a4" else LETTER
    doc = SimpleDocTemplate(
        output, pagesize=pagesize, rightMargin=16 * mm, leftMargin=16 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm, title=title,
        author="Opportunity Finder",
    )
    styles = getSampleStyleSheet()
    regular_font, bold_font = _pdf_fonts(title + " " + text[:5000])
    styles.add(ParagraphStyle(
        name="SourceTitle", parent=styles["Title"], fontName=bold_font,
        fontSize=16, leading=19, alignment=TA_CENTER, spaceAfter=10,
    ))
    styles.add(ParagraphStyle(
        name="SourceMeta", parent=styles["Normal"], fontName=regular_font,
        fontSize=8.5, leading=11, textColor="#555555", spaceAfter=8,
    ))
    styles["BodyText"].fontName = regular_font
    styles["BodyText"].fontSize = 9.5
    styles["BodyText"].leading = 13
    styles["BodyText"].spaceAfter = 7
    story = [
        Paragraph(_xml(title), styles["SourceTitle"]),
        Paragraph(_xml(f"Source: {url}") + "<br/>" + _xml(f"Retrieved: {retrieved_at}"),
                  styles["SourceMeta"]),
    ]
    for block in re.split(r"\n\s*\n", text or ""):
        block = " ".join(block.split())
        for chunk in textwrap.wrap(block, width=4500, break_long_words=False,
                                   break_on_hyphens=False) if block else []:
            story.append(Paragraph(_xml(chunk), styles["BodyText"]))
    doc.build(story)
    return output.getvalue()


def _xml(value: str) -> str:
    return (value or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def page_system_for(opportunity: dict) -> str:
    country = (opportunity.get("country") or opportunity.get("location") or "").lower()
    return "letter" if any(x in country for x in ("united states", "usa", "canada")) else "a4"
