"""Integration checks for account isolation and the commercial product flow."""

from __future__ import annotations

import base64
import io
import os
import unittest
import json
from unittest.mock import patch

os.environ.setdefault("SECRET_KEY", "integration-test-secret-not-for-production")

from app import app
from finder import pipeline, store
from finder import application_pack
from finder.matching import score_for_profile
from finder.profiles import build_profile
from finder.industry import _direct_job_result, _parse_nhs_xml, _worldwide_job_result
from finder.docgen import tailor_cv_docx, webpage_pdf
from finder.enrich import explicit_deadline
from finder.pack_requirements import detect_requirements

EMAILS = ("product-test-one@example.invalid", "product-test-two@example.invalid")
OPP_IDS = ("__product_academic__", "__product_industry__")


def cleanup_test_records() -> None:
    conn = store.connect()
    try:
        users = conn.execute("SELECT id FROM users WHERE email IN (?,?)", EMAILS).fetchall()
        for user in users:
            uid = user["id"]
            conn.execute("DELETE FROM application_pack_files WHERE user_id=?", (uid,))
            conn.execute("DELETE FROM application_packs WHERE user_id=?", (uid,))
            conn.execute("DELETE FROM provider_configs WHERE user_id=?", (uid,))
            conn.execute("DELETE FROM user_matches WHERE user_id=?", (uid,))
            conn.execute("DELETE FROM user_profiles WHERE user_id=?", (uid,))
            conn.execute("DELETE FROM users WHERE id=?", (uid,))
        for opp_id in OPP_IDS:
            conn.execute("DELETE FROM opportunities WHERE id=?", (opp_id,))
        conn.commit()
    finally:
        conn.close()


def sample(opp_id: str, track: str) -> dict:
    return {
        "id": opp_id, "source": "Product self test", "source_key": "producttest",
        "title": ("Assistant Professor of Construction Management" if track == "academic"
                  else "Senior Infrastructure Project Manager"),
        "org": "Example Engineering Group", "department": "Civil and Infrastructure",
        "location": "Toronto, Canada", "country": "Canada", "country_tier": "target",
        "url": f"https://example.invalid/{opp_id}",
        "description": "Construction management, water infrastructure, digital twins, BIM and project delivery.",
        "posted": "2026-09-20", "deadline": "2026-12-31", "days_left": 100,
        "role_key": "assistant" if track == "academic" else "industry",
        "role_label": "Assistant Professor / Lecturer" if track == "academic" else "Industry role",
        "score": 50, "matched_terms": [], "flags": [], "positives": [], "breakdown": {},
        "query": "product test", "enriched": 1, "career_track": track, "status": "new",
    }


class ProductFlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config.update(TESTING=True, SECRET_KEY=os.environ["SECRET_KEY"])
        cleanup_test_records()
        conn = store.connect()
        store.upsert_many(conn, [sample(OPP_IDS[0], "academic"), sample(OPP_IDS[1], "industry")])
        conn.close()

    @classmethod
    def tearDownClass(cls):
        cleanup_test_records()

    def _register(self, email: str):
        client = app.test_client()
        client.get("/register")
        with client.session_transaction() as session:
            csrf = session["csrf"]
        response = client.post("/register", data={"name": "Product Test User", "email": email,
                               "password": "A-long-test-password-42", "csrf_token": csrf})
        self.assertEqual(response.status_code, 302)
        with client.session_transaction() as session:
            return client, session["csrf"], session["user_id"]

    def _upload(self, client, csrf: str, kind: str, text: str):
        response = client.post(f"/api/profiles/{kind}",
            data={"cv": (io.BytesIO(text.encode()), f"{kind}.txt")},
            headers={"X-CSRF-Token": csrf}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()

    def test_occupational_domain_gate_blocks_incidental_health_words(self):
        cv = ("Registered Nurse with BSc Nursing. Staff nurse in intensive care, "
              "clinical nursing, medication administration and ward practice. " * 10)
        structured = build_profile(cv, "industry")
        self.assertEqual(structured["primary_domain"], "nursing")
        profile = {"profile": structured, "cv_text": cv}
        base = {"org": "Example", "department": "", "location": "London",
                "country_tier": "target", "days_left": 30}
        nurse_score = score_for_profile({**base, "title": "Registered Nurse - ICU",
            "description": "Provide nursing care to patients in intensive care."}, profile)[0]
        data_score = score_for_profile({**base, "title": "Data Analyst - Luxury Platform",
            "description": "Analyse customer health and care data in dashboards."}, profile)[0]
        care_score = score_for_profile({**base, "title": "Customer Care Manager",
            "description": "Lead customer care and wellbeing services."}, profile)[0]
        self.assertGreaterEqual(nurse_score, 70)
        self.assertLessEqual(data_score, 18)
        self.assertLessEqual(care_score, 18)

    def test_free_search_provider_is_the_default(self):
        from finder.providers import public_catalog
        self.assertEqual(public_catalog()["search"][0], "duckduckgo")

    def test_official_nhs_nursing_feed_is_normalized(self):
        payload = b"""<nhsJobs><vacancyDetails><closeDate>2026-10-20</closeDate>
        <description>Lead a theatre team and provide safe patient services.</description>
        <employer>Example NHS Trust</employer><id>123</id><locations>
        <locations>Manchester, M1 1AA</locations></locations>
        <postDate>2026-09-21T08:00:00</postDate><title>Theatre Team Leader</title>
        <url>https://www.jobs.nhs.uk/candidate/jobadvert/123</url></vacancyDetails></nhsJobs>"""
        rows = _parse_nhs_xml(payload)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["source_key"], "nhs_jobs")
        self.assertEqual(rows[0]["country"], "United Kingdom")
        self.assertEqual(rows[0]["deadline"], "2026-10-20")
        self.assertIn("nursing", rows[0]["description"].lower())

    def test_direct_employer_job_search_discards_landing_pages(self):
        direct = _direct_job_result({
            "title": "Next Health - Registered Nurse - Los Angeles - Lever",
            "url": "https://jobs.lever.co/next-health/2d9420fe-3b3e-4ba9-b4a8-1621bf8527fa/apply",
            "snippet": "Registered Nurse position in Los Angeles.",
        }, "nursing", "registered nurse")
        self.assertIsNotNone(direct)
        self.assertEqual(direct["source_key"], "web_lever")
        self.assertFalse(direct["url"].endswith("/apply"))
        landing = _direct_job_result({
            "title": "Nursing jobs", "url": "https://jobs.lever.co/next-health",
            "snippet": ""}, "nursing", "registered nurse")
        self.assertIsNone(landing)

    def test_worldwide_search_keeps_direct_jobs_not_articles(self):
        direct = _worldwide_job_result({
            "title": "Registered Nurse - Emergency Department | Example Health",
            "url": "https://careers.example.ca/jobs/registered-nurse-123",
            "snippet": "Apply for this registered nurse role.",
        }, "registered nurse", "Canada", "serper", "registered nurse Canada")
        self.assertIsNotNone(direct)
        self.assertEqual(direct["country"], "Canada")
        article = _worldwide_job_result({
            "title": "How to become a registered nurse in Canada",
            "url": "https://example.ca/blog/nursing-guide", "snippet": "Guide",
        }, "registered nurse", "Canada", "serper", "registered nurse Canada")
        self.assertIsNone(article)

    def test_unknown_profession_still_gets_search_titles(self):
        cv = ("PROFESSIONAL EXPERIENCE\nCybersecurity Consultant\n"
              "Protected enterprise networks and conducted security reviews.\n" * 8)
        profile = build_profile(cv, "industry")
        self.assertIn("Cybersecurity Consultant", profile["target_titles"])

    def test_application_requirements_detect_diversity_and_word_limits(self):
        material = (
            "Submit a curriculum vitae, cover letter, research statement, teaching statement, "
            "and a diversity statement limited to 750 words."
        )
        result = detect_requirements(material, "academic")
        items = {item["key"]: item for item in result["submission_items"]}
        self.assertIn("diversity_statement", items)
        self.assertEqual(items["diversity_statement"]["limit"], "750 words")
        self.assertIn("research_statement", result["documents"])

    def test_advert_application_window_overrides_aggregator_deadline(self):
        text = ("Review of applications will start from July 2026 and continue until "
                "August 31, 2026 or until the post is filled.")
        self.assertEqual(explicit_deadline(text), "2026-08-31")

    def test_docx_cv_tailoring_preserves_unedited_content_and_layout(self):
        from docx import Document
        from docx.shared import Inches
        source = Document()
        source.sections[0].left_margin = Inches(0.55)
        source.add_heading("Curriculum Vitae", 0)
        source.add_paragraph("Infrastructure engineer and researcher")
        source.add_paragraph("Managed water infrastructure research projects", style="List Bullet")
        source.add_paragraph("Publication record remains unchanged")
        stream = io.BytesIO(); source.save(stream)
        tailored, audit = tailor_cv_docx(stream.getvalue(), {
            "replacements": [{"find": "Infrastructure engineer and researcher",
                              "replacement": "Civil infrastructure engineer and researcher"}],
            "additions": [{"after": "Managed water infrastructure research projects",
                           "text": "Applied project management methods to infrastructure research",
                           "style": "List Bullet"}],
        })
        result = Document(io.BytesIO(tailored))
        text = "\n".join(p.text for p in result.paragraphs)
        self.assertIn("Civil infrastructure engineer and researcher", text)
        self.assertIn("Publication record remains unchanged", text)
        self.assertIn("Applied project management methods", text)
        self.assertAlmostEqual(result.sections[0].left_margin.inches, 0.55, places=2)
        self.assertEqual(audit["replacements_applied"], 1)
        self.assertEqual(audit["additions_applied"], 1)

    def test_web_research_snapshot_is_a_valid_pdf(self):
        data = webpage_pdf("Department Modules", "https://example.invalid/modules",
                           "Module A covers infrastructure management.\n\nModule B covers BIM.",
                           "2026-09-23T00:00:00+00:00")
        self.assertTrue(data.startswith(b"%PDF"))
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        self.assertGreaterEqual(len(reader.pages), 1)
        self.assertIn("infrastructure management", reader.pages[0].extract_text().lower())

    def test_refresh_button_selects_only_the_current_track(self):
        client = app.test_client()
        was_running = pipeline.STATE["running"]
        pipeline.STATE["running"] = False
        try:
            with patch.object(pipeline, "run_refresh_background") as start:
                response = client.post("/api/refresh?track=industry")
                self.assertEqual(response.status_code, 200)
                start.assert_called_once_with({"industry"})
        finally:
            pipeline.STATE["running"] = was_running

    def test_end_to_end_account_isolation(self):
        one, csrf_one, uid_one = self._register(EMAILS[0])
        two, csrf_two, uid_two = self._register(EMAILS[1])
        cv_one = ("Civil engineer and construction management researcher. Water infrastructure, "
                  "digital twin, BIM, project delivery and asset management. " * 12)
        cv_two = ("Molecular biology laboratory specialist in genomics, proteins, cells and clinical assays. " * 14)
        self._upload(one, csrf_one, "academic", cv_one)
        self._upload(one, csrf_one, "industry", cv_one)
        self._upload(two, csrf_two, "academic", cv_two)

        preferences = one.post("/api/profiles/industry/preferences", json={
            "target_titles": "Infrastructure Project Manager, Civil Engineer",
            "preferred_locations": "Canada, Australia",
        }, headers={"X-CSRF-Token": csrf_one})
        self.assertEqual(preferences.status_code, 200, preferences.get_data(as_text=True))
        self.assertEqual(preferences.get_json()["profile"]["preferred_locations"],
                         ["Canada", "Australia"])

        self.assertEqual(one.get("/profile").status_code, 200)
        self.assertEqual(one.get("/settings").status_code, 200)
        was_running = pipeline.STATE["running"]
        pipeline.STATE["running"] = True
        try:
            dashboard = one.get("/?track=academic")
            self.assertEqual(dashboard.status_code, 200, dashboard.get_data(as_text=True))
            self.assertIn(b"Academic jobs", dashboard.data)
            self.assertIn(b"Industry jobs", dashboard.data)
        finally:
            pipeline.STATE["running"] = was_running

        query = "/api/opportunities?track=academic&min_score=0&limit=100&q=Example%20Engineering%20Group"
        first = one.get(query).get_json()
        second = two.get(query).get_json()
        item_one = next(x for x in first["items"] if x["id"] == OPP_IDS[0])
        item_two = next(x for x in second["items"] if x["id"] == OPP_IDS[0])
        self.assertGreater(item_one["score"], item_two["score"])

        missing_csrf = one.post("/api/status", json={"id": OPP_IDS[0], "status": "saved",
                                                       "track": "academic"})
        self.assertEqual(missing_csrf.status_code, 400)
        saved = one.post("/api/status", json={"id": OPP_IDS[0], "status": "saved",
                                              "track": "academic"},
                         headers={"X-CSRF-Token": csrf_one})
        self.assertEqual(saved.status_code, 200)
        other_view = two.get(f"/api/opportunity/{OPP_IDS[0]}?track=academic").get_json()
        self.assertEqual(other_view["status"], "new")

        config = one.post("/api/providers", json={
            "ai_provider": "poe", "ai_model": "Claude-Sonnet-4.6",
            "ai_base_url": "https://api.poe.com/v1/chat/completions",
            "ai_key": "plaintext-test-key-must-not-remain",
            "search_provider": "duckduckgo", "search_key": "",
        }, headers={"X-CSRF-Token": csrf_one})
        self.assertEqual(config.status_code, 200, config.get_data(as_text=True))
        conn = store.connect()
        try:
            stored = conn.execute("SELECT ai_key_enc FROM provider_configs WHERE user_id=?",
                                  (uid_one,)).fetchone()["ai_key_enc"]
            self.assertNotIn("plaintext-test-key", stored)
            conn.execute("""INSERT INTO application_packs
                (id,user_id,opportunity_id,profile_kind,status,zip_blob,created_at)
                VALUES (?,?,?,?,?,?,?)""",
                ("__pack_test__", uid_one, OPP_IDS[0], "academic", "ready",
                 base64.b64encode(b"PK-test-private").decode(), "2026-09-21"))
            conn.commit()
        finally:
            conn.close()
        self.assertEqual(two.get("/api/packs/__pack_test__/download").status_code, 404)
        own_download = one.get("/api/packs/__pack_test__/download")
        self.assertEqual(own_download.status_code, 200)
        self.assertIn("Assistant Professor of Construction Management at Example Engineering Group.zip",
                      own_download.headers.get("Content-Disposition", ""))
        self.assertEqual(two.get("/applications/__pack_test__").status_code, 404)
        self.assertEqual(one.get("/applications/__pack_test__").status_code, 200)
        self.assertIn(b"Application workspace", one.get("/applications/__pack_test__").data)

        conn = store.connect()
        try:
            conn.execute("""INSERT INTO application_packs
                (id,user_id,opportunity_id,profile_kind,status,created_at)
                VALUES (?,?,?,?,?,?)""",
                ("__generated_pack__", uid_one, OPP_IDS[0], "academic", "working", "2026-09-21"))
            conn.commit()
        finally:
            conn.close()
        plan = json.dumps({
            "submission_items": [
                {"key": "cover_letter", "required": True, "detail": "Requested", "limit": ""},
                {"key": "tailored_cv", "required": True, "detail": "Requested", "limit": ""},
            ],
            "criteria": [{"criterion": "Construction management expertise",
                          "priority": "Essential", "evidence": "CV states this expertise"}],
            "gaps": [], "cv_edits": {"replacements": [], "additions": []},
        })
        generated = "## Application\n\n" + (
            "Evidence-based tailored cover letter content grounded in the curriculum vitae. " * 60
        )
        sources = [{"title": "Department courses", "url": "https://example.invalid/courses",
                    "snippet": "Construction management course information.",
                    "text": "Construction management course information.",
                    "kind": "web page", "query": "department courses",
                    "retrieved_at": "2026-09-21"}]
        research = {"sources": sources, "attachments": [],
                    "advert_text": sample(OPP_IDS[0], "academic")["description"]}
        with patch.object(application_pack, "chat", side_effect=[plan, generated]), \
             patch.object(application_pack, "collect_application_research", return_value=research):
            application_pack._generate("__generated_pack__", uid_one, OPP_IDS[0], "academic",
                                       ["cover_letter", "tailored_cv"])
        pack = application_pack.get_pack(uid_one, "__generated_pack__", include_blob=True)
        self.assertEqual(pack["status"], "ready", pack.get("error"))
        self.assertEqual(pack["progress"], 100)
        self.assertIsNone(pack["zip_blob"])
        files = application_pack.list_pack_files(uid_one, "__generated_pack__")
        names = {item["filename"] for item in files}
        self.assertIn("Application Documents/Cover Letter - Example Engineering Group.docx", names)
        self.assertTrue(any(name.startswith("CV/Tailored CV Working Copy") for name in names))
        self.assertTrue({"Research/Research Sources and Links.docx",
                         "Job Materials/Job Advertisement Snapshot.docx",
                         "Job Materials/Job Advertisement Snapshot.pdf",
                         "Application Requirements and Evidence Plan.docx",
                         "manifest.json", "README FIRST.txt"}.issubset(names))
        status = one.get("/api/packs/__generated_pack__")
        self.assertEqual(status.status_code, 200)
        payload = status.get_json()
        self.assertEqual(payload["status"], "ready")
        self.assertEqual(payload["progress"], 100)
        self.assertEqual(len(payload["steps"]), 5)
        self.assertTrue(all(step["status"] == "complete" for step in payload["steps"]))
        cover = next(item for item in files if item["filename"].startswith("Application Documents/"))
        own_file = one.get(f"/api/packs/__generated_pack__/files/{cover['id']}")
        self.assertEqual(own_file.status_code, 200)
        self.assertTrue(own_file.data.startswith(b"PK"))
        self.assertEqual(two.get(f"/api/packs/__generated_pack__/files/{cover['id']}").status_code, 404)
        self.assertIn(b"Download only the files you need",
                      one.get("/applications/__generated_pack__").data)


if __name__ == "__main__":
    unittest.main(verbosity=2)
