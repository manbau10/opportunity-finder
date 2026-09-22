"""Print compact structural summaries for DOCX and PDF application files."""

from __future__ import annotations

import argparse
import collections
import pathlib

from docx import Document
from pypdf import PdfReader


def summarize_docx(path: pathlib.Path) -> None:
    doc = Document(path)
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    words = sum(len(p.text.split()) for p in paragraphs)
    styles = collections.Counter(p.style.name for p in paragraphs)
    fonts = collections.Counter(
        (run.font.name or "(inherited)", round(run.font.size.pt, 1) if run.font.size else None)
        for p in paragraphs for run in p.runs if run.text.strip()
    )
    headings = [p.text.strip() for p in paragraphs if p.style.name.startswith(("Title", "Heading"))]
    sections = [
        {
            "size_in": (round(s.page_width.inches, 2), round(s.page_height.inches, 2)),
            "margins_in": (round(s.top_margin.inches, 2), round(s.right_margin.inches, 2),
                           round(s.bottom_margin.inches, 2), round(s.left_margin.inches, 2)),
        }
        for s in doc.sections
    ]
    print(f"DOCX\t{path}\n  paragraphs={len(paragraphs)} words={words} tables={len(doc.tables)} sections={sections}")
    print(f"  styles={styles.most_common(8)} fonts={fonts.most_common(8)}")
    print(f"  headings={headings[:20]}")
    print("  opening=" + " | ".join(p.text.strip().replace("\n", " ")[:180] for p in paragraphs[:8]))


def summarize_pdf(path: pathlib.Path) -> None:
    reader = PdfReader(path)
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    clean = [" ".join(x.split()) for x in text.splitlines() if x.strip()]
    print(f"PDF\t{path}\n  pages={len(reader.pages)} words={len(text.split())}")
    print("  opening=" + " | ".join(clean[:12])[:1200])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+")
    args = parser.parse_args()
    for raw in args.paths:
        root = pathlib.Path(raw)
        files = [root] if root.is_file() else sorted(root.glob("*"))
        for path in files:
            try:
                if path.suffix.lower() == ".docx":
                    summarize_docx(path)
                elif path.suffix.lower() == ".pdf":
                    summarize_pdf(path)
            except Exception as exc:
                print(f"ERROR\t{path}\t{type(exc).__name__}: {exc}")


if __name__ == "__main__":
    main()
